from ..pdf_as_json_model import PrecisionPDFModel, MinerPDFModel
from ..ms_model import MSWordModel
from typing import Dict, List
import docx
from docx.enum.style import WD_STYLE_TYPE
import base64
from io import BytesIO

# TODO: переделать на абстрактный класс
class PrecisionPDFToMSWord:
    def convert(self, input_model: PrecisionPDFModel, output_model: MSWordModel)-> None:
        pages: List[Dict] = input_model.pdf_json["pages"]
        output_model.doc = docx.Document()
        styles = output_model.doc.styles
        self.LABELS = {
            "text": "Normal",
            "header": "Heading 1",
            "list": "List",
            "table": "Normal",
            # "figure": "Normal",
        }
        for page in pages:
            self.add_page(output_model.doc, page, styles)
      

    def add_page(self, doc: docx.Document, page: Dict, styles)-> None:
        for regions in page["regions"]:
            label = regions["label"]
            text = regions["text"]
            if label == "figure" or regions.get("base64"):
                self._add_figure(doc, regions)
                continue
            doc.add_paragraph(text).style = styles[self.LABELS[label]]
        doc.add_page_break()

    def _add_figure(self, doc: docx.Document, region: Dict) -> None:
        b64 = region.get("base64")
        if not b64:
            return
        if str(b64).startswith("data:"):
            b64 = str(b64).split(",", 1)[1]
        try:
            stream = BytesIO(base64.b64decode(b64))
            shape = doc.add_picture(stream)
            section = doc.sections[0]
            usable_width = section.page_width - section.left_margin - section.right_margin
            if shape.width > usable_width:
                ratio = usable_width / shape.width
                shape.width = usable_width
                shape.height = int(shape.height * ratio)
        except Exception:
            pass

class MinerPDFToMSWord:
    def convert(self, input_model: MinerPDFModel, output_model: MSWordModel)-> None:
        pages: List[Dict] = input_model.pdf_json["pages"]
        output_model.doc = docx.Document()
        styles = output_model.doc.styles
        self.LABELS = {
            "text": "Normal",
            "header": "Heading 1",
            "list": "List",
            "table": "Normal",
            "other": "Normal",
            # "figure": "Normal",
        }
        for page in pages:
            self.add_page(output_model.doc, page, styles)
      

    def add_page(self, doc: docx.Document, page: Dict, styles)-> None:
        for regions in page["regions"]:
            label = regions["label"]
            text = regions["text"]
            if label == "figure" or regions.get("base64"):
                self._add_figure(doc, regions)
                continue
            doc.add_paragraph(text).style = styles[self.LABELS[label]]
        doc.add_page_break()

    def _add_figure(self, doc: docx.Document, region: Dict) -> None:
        b64 = region.get("base64")
        if not b64:
            return
        if str(b64).startswith("data:"):
            b64 = str(b64).split(",", 1)[1]
        try:
            stream = BytesIO(base64.b64decode(b64))
            shape = doc.add_picture(stream)
            section = doc.sections[0]
            usable_width = section.page_width - section.left_margin - section.right_margin
            if shape.width > usable_width:
                ratio = usable_width / shape.width
                shape.width = usable_width
                shape.height = int(shape.height * ratio)
        except Exception:
            pass